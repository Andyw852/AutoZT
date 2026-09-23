from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'\\wsl.localhost\Ubuntu\home\user\software\AutoZT\docs\AutoZT_workflow_software_comparison.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color='D9D9D9', size='6'):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = b.find(qn('w:'+edge))
        if el is None:
            el = OxmlElement('w:'+edge); b.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), size); el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)

def margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e = mar.find(qn('w:'+k))
        if e is None: e = OxmlElement('w:'+k); mar.append(e)
        e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa')

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

def set_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr(); w = tcPr.find(qn('w:tcW'))
            if w is None: w = OxmlElement('w:tcW'); tcPr.append(w)
            w.set(qn('w:w'), str(int(width*1440))); w.set(qn('w:type'), 'dxa')

def cell_text(cell, text, bold=False, size=8.2, color='000000', align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]; p.alignment = align; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text); r.bold = bold; r.font.name = 'Times New Roman'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun'); r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)

def add_table(doc, headers, rows, widths, font_size=8.2):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    set_widths(t, widths); borders(t); set_repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        shade(t.rows[0].cells[i], '1F4E78'); cell_text(t.rows[0].cells[i], h, True, font_size, 'FFFFFF', WD_ALIGN_PARAGRAPH.CENTER)
    for ri,row in enumerate(rows):
        cells = t.add_row().cells
        for i,val in enumerate(row):
            if ri % 2 == 1: shade(cells[i], 'EAF2F8')
            cell_text(cells[i], val, False, font_size, '000000', WD_ALIGN_PARAGRAPH.LEFT if i not in (0,1) else WD_ALIGN_PARAGRAPH.CENTER)
    return t

doc = Document()
sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width = Inches(11); sec.page_height = Inches(8.5)
sec.left_margin = Inches(.7); sec.right_margin = Inches(.7); sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
styles = doc.styles
styles['Normal'].font.name = 'Times New Roman'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun'); styles['Normal'].font.size = Pt(9)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
r = p.add_run('AutoZT 与同类材料计算工作流软件对比'); r.bold = True; r.font.name = 'Times New Roman'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun'); r.font.size = Pt(16); r.font.color.rgb = RGBColor(0,0,0)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8); p.paragraph_format.line_spacing = 1.05
r = p.add_run('本表只保留与 AutoZT 有直接可比性的系统：材料原生平台、机器学习势流水线、原子模拟接口和通用科学工作流引擎。优缺点按公开定位与 AutoZT 当前实现归纳；“未见专用支持”表示资料中未发现面向该功能的原生流程，不等于无法通过二次开发实现。'); r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'SimSun'); r.font.size=Pt(9)

doc.add_paragraph('表 1  核心能力与工程取舍').runs[0].bold = True
headers = ['系统', '定位', '材料工作流与组件', '后端与集群', '扩展 / 智能体接口', '主要优点与短板']
rows = [
['pyiron', '材料科学 IDE / 作业平台', '原子结构、DFT、MD、势函数；材料组件完整', '可接 VASP、QE、LAMMPS、ASE；支持 HPC', 'Python Job / Project；无原生 LLM 控制', '优：材料生态成熟、交互和批处理兼顾。短：热电专用闭环、物理验收闸和 LLM 风险网关需自行实现。'],
['Simmate', '材料数据库与高通量工作流', '结构、性质计算、数据库和 Web 工作流', '支持多种 DFT 代码；可部署到 HPC', 'Workflow / App / 数据模型；无原生 LLM 控制', '优：适合高通量和结果入库。短：MACE、随机位移热输运及跨集群故障语义不是核心能力。'],
['AFLOW / AFLOWπ', '高通量材料发现平台', '标准化 DFT 输入、枚举、性质计算和数据库', '面向 HPC 高通量；以 DFT 为主', '脚本与模板扩展；无原生 LLM 控制', '优：规模化材料筛选能力强。短：偏数据库和 DFT 生产，难直接覆盖 MACE/热电全链与人工审批。'],
['MatFlow', '可复现材料模拟工作流', '配置化流程、模板、参数和结果记录', '可接 HPC、容器和外部程序', 'YAML / 模板 / Workflow；无原生 LLM 控制', '优：流程复现和参数管理清晰。短：材料计算组件与热输运判据需要用户组装。'],
['DP-GEN', '机器学习势主动学习流水线', '探索、标注、训练、验证和迭代数据生成', '适合 HPC 批量任务；以 ML 势训练为中心', '阶段配置和脚本；不是通用 LLM 工作流', '优：ML 势数据闭环成熟。短：不是通用材料生产平台，VASP/MACE 热电流程需另行编排。'],
['ASE', '原子模拟统一 Python 接口', '结构操作、计算器、优化、采样和分析', '由用户配置计算器和调度器', 'Calculator / Python 扩展；无原生状态监督', '优：后端替换灵活，适合搭积木。短：任务依赖、失败分类、审计和生产监控不由 ASE 负责。'],
['Parsl / Pegasus / Nextflow', '通用科学工作流编排', 'DAG、并行任务、重试、容器或数据流', '多集群、云和调度器适配较强', 'Python / DSL；无材料语义和原生 LLM', '优：通用并行与跨站点执行成熟。短：VASP/MACE 输入、物理判据、热输运验收需自行开发。'],
['AutoZT', '面向热电材料的生产执行与监督', 'VASP / MACE 技能、MLFF、电子与晶格热输运、ZT；物理判据驱动 DAG', 'jzzn CPU、A800 GPU、3090 无 SLURM；按材料和步骤混合调度', 'skill 目录可改；MCP 工具、风险分级、人工批准；稳态推进不依赖 LLM', '优：热电专用、跨 DFT/MLFF、失败可分类恢复、逐文件哈希与会话审计。短：材料数据库和开放式发现能力不以此为目标，技能范围需持续扩展。'],
]
add_table(doc, headers, rows, [1.25,1.35,2.25,1.75,1.65,2.05], 7.7)

doc.add_paragraph().paragraph_format.space_after = Pt(2)
doc.add_paragraph('表 2  选型建议').runs[0].bold = True
headers2 = ['需求', '优先考虑', '原因与边界']
rows2 = [
['材料数据库与高通量筛选', 'Simmate / AFLOW', '数据库、结构枚举和高通量能力更成熟；需要另接 AutoZT 的热输运和运维判据。'],
['交互式材料计算与多代码研究', 'pyiron', '适合研究者快速组合 VASP、QE、LAMMPS 和 ASE；生产级热电闭环仍需二次开发。'],
['机器学习势主动学习', 'DP-GEN + AutoZT mlff', 'DP-GEN 擅长训练数据闭环；AutoZT 擅长把 MLFF 与 VASP/MACE 热输运步骤接起来。'],
['自建新材料工作流', 'ASE + Parsl / Nextflow', '组件和调度自由度最高，但物理验收、失败恢复、审计和安全边界由项目自行承担。'],
['热电材料批量生产与监督', 'AutoZT', '当前唯一以电子输运、晶格热输运、厚度/二维修正、MLFF 和多集群运行治理为一体的方案。'],
]
add_table(doc, headers2, rows2, [2.2,2.0,6.1], 8.2)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
r=p.add_run('说明：AiiDA、FireWorks、atomate2 和 Taskflow 已在原对比中列出，本表不重复展开；它们可作为 AutoZT 的通用工作流基线。公开资料检索截至 2026 年 9 月，具体版本和后端支持应在部署前按官方文档复核。'); r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'SimSun'); r.font.size=Pt(8); r.italic=True

doc.save(OUT)
print(OUT)
